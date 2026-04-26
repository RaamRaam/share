"""
scripts/create_sample_excel.py — Generate tests/test_cases.xlsx with sample data.

Run once to get a working template, then edit it with your real test cases.

  python scripts/create_sample_excel.py
  python scripts/create_sample_excel.py tests/my_tests.xlsx
"""

import json
import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

OUTPUT = Path("tests/test_cases.xlsx")


def p(d: dict) -> str:
    """Serialize a dict to compact JSON for the params cell."""
    return json.dumps(d, ensure_ascii=False)


HEADERS = [
    "Test Case ID",
    "Test Scenario",
    "Step Number",
    "Step Function",
    "Step Parameters",
    "Module",
    "Step Description",
    "Expected Result",
]

# ── Sample rows (TC ID, Scenario, Step, Function, Params dict, Module, Desc, Expected)
ROWS = [
    # ── TC001: Open app and verify home page ─────────────────────────────────
    ("TC001", "Open app and verify home page", 1,
     "go_to",        p({"url": "${APP_URL}"}),
     "home",         "Navigate to the application URL", ""),
    ("TC001", "", 2,
     "check_url",    p({"expected": "${APP_URL}"}),
     "",             "Assert URL loaded correctly", ""),
    ("TC001", "", 3,
     "check_visible",p({"selector": "header, h1"}),
     "",             "Assert page header is visible", ""),

    # ── TC002: Login via SSO ──────────────────────────────────────────────────
    ("TC002", "Login via SSO and land on dashboard", 1,
     "go_to",        p({"url": "${APP_URL}"}),
     "auth",         "Open application", ""),
    ("TC002", "", 2,
     "click",        p({"text": "Sign In", "exact": False}),
     "auth",         "Click the Sign In button", ""),
    ("TC002", "", 3,
     "wait_for_url", p({"pattern": "**login.microsoftonline.com**", "timeout": 15000}),
     "auth",         "Wait for Microsoft login page", ""),
    ("TC002", "", 4,
     "click",        p({"text": "${SSO_EMAIL}", "exact": False}),
     "auth",         "Click SSO account tile", ""),
    ("TC002", "", 5,
     "wait_for_url", p({"pattern": "**/dashboard**", "timeout": 20000}),
     "auth",         "Wait for dashboard to load", "/dashboard"),

    # ── TC003: Fill form by placeholder and label ─────────────────────────────
    ("TC003", "Fill new-record form and submit", 1,
     "click",        p({"text": "New Record"}),
     "records",      "Open the New Record form", ""),
    ("TC003", "", 2,
     "fill_input",   p({"placeholder": "Enter title", "value": "Automation Test Record"}),
     "records",      "Fill Title field by placeholder", ""),
    ("TC003", "", 3,
     "fill_input",   p({"label": "Description", "value": "Created by automation"}),
     "records",      "Fill Description field by label text", ""),
    ("TC003", "", 4,
     "fill_input",   p({"placeholder": "Owner email", "value": "${SSO_EMAIL}"}),
     "records",      "Fill Owner with SSO email from env", ""),
    ("TC003", "", 5,
     "click",        p({"text": "Submit", "validate_url": "/records"}),
     "records",      "Submit form and assert redirect", "/records"),

    # ── TC004: Search and capture adjacent text ───────────────────────────────
    ("TC004", "Search for a record and read its status", 1,
     "search_input", p({"placeholder": "Search records…", "value": "Automation Test"}),
     "records",      "Search for the new record", ""),
    ("TC004", "", 2,
     "get_adjacent_text", p({"label": "Status"}),
     "records",      "Read Status value next to its label", "Active"),
    ("TC004", "", 3,
     "get_adjacent_text", p({"label": "Owner"}),
     "records",      "Read Owner value next to its label", ""),

    # ── TC005: Table row icon click ───────────────────────────────────────────
    ("TC005", "Edit a record via table action icon", 1,
     "click",        p({"text": "Records"}),
     "records",      "Go to Records list page", ""),
    ("TC005", "", 2,
     "click_row_icon", p({"search_text": "Automation Test Record",
                          "icon_index": 0}),
     "records",      "Click the first icon (Edit) in the matching row", ""),
    ("TC005", "", 3,
     "wait_for_page",p({}),
     "records",      "Wait for edit page to load", ""),

    # ── TC006: Table row icon click by name ───────────────────────────────────
    ("TC006", "Delete a record via named icon in table", 1,
     "click_row_icon", p({"search_text": "Automation Test Record",
                          "icon_name": "delete"}),
     "records",      "Click the Delete icon by aria-label", ""),
    ("TC006", "", 2,
     "click",        p({"text": "Confirm"}),
     "records",      "Confirm deletion in the dialog", ""),

    # ── TC007: Demo mouse sweep ───────────────────────────────────────────────
    ("TC007", "Demo: sweep cursor to show automation", 1,
     "go_to",        p({"url": "${APP_URL}"}),
     "",             "Open the app", ""),
    ("TC007", "", 2,
     "demo_mouse_move", p({"steps": 15, "duration_ms": 2000}),
     "",             "Sweep cursor across the page", ""),
]

# Column widths (characters)
COL_WIDTHS = [12, 35, 10, 20, 55, 12, 40, 20]


def create_excel(path: Path = OUTPUT) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Test Cases"

    # ── Header row ────────────────────────────────────────────────────────────
    header_fill = PatternFill("solid", fgColor="2E4057")
    header_font = Font(bold=True, color="FFFFFF", size=11)
    for col, (header, width) in enumerate(zip(HEADERS, COL_WIDTHS), start=1):
        cell              = ws.cell(row=1, column=col, value=header)
        cell.fill         = header_fill
        cell.font         = header_font
        cell.alignment    = Alignment(horizontal="center", vertical="center",
                                      wrap_text=True)
        ws.column_dimensions[get_column_letter(col)].width = width

    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "A2"

    # ── Data rows ─────────────────────────────────────────────────────────────
    alt_fill = PatternFill("solid", fgColor="F0F4F8")   # light blue-grey
    cur_tc   = None
    shade    = False

    for row_num, row_data in enumerate(ROWS, start=2):
        tc_id = row_data[0]
        if tc_id != cur_tc:
            cur_tc = tc_id
            shade  = not shade

        fill = alt_fill if shade else None

        for col, value in enumerate(row_data, start=1):
            cell           = ws.cell(row=row_num, column=col, value=value)
            cell.alignment = Alignment(vertical="top", wrap_text=(col == 5))
            if fill:
                cell.fill = fill

        # Bold TC ID
        ws.cell(row=row_num, column=1).font = Font(bold=bool(row_data[0]))

    # ── Instruction sheet ─────────────────────────────────────────────────────
    info = wb.create_sheet("Instructions")
    instructions = [
        ("Column",          "Description",                              "Example"),
        ("A  Test Case ID", "Unique ID for the test case",              "TC001"),
        ("B  Test Scenario","Human-readable scenario (first row only)", "Login Flow"),
        ("C  Step Number",  "Sequential step number",                   "1"),
        ("D  Step Function","Action function name (see list below)",    "click"),
        ("E  Step Parameters", 'JSON object of keyword args',           '{"text": "Save"}'),
        ("F  Module",       "Optional — URL path segment or label for override lookup",
                            "dashboard"),
        ("G  Step Description", "Optional — shown in reports",          "Click Save button"),
        ("H  Expected Result",  "Optional — stored for reference",      "/dashboard"),
        ("", "", ""),
        ("AVAILABLE FUNCTIONS", "", ""),
        ("go_to",            "Navigate to URL",                         '{"url": "${APP_URL}"}'),
        ("click",            "Click element by text",                   '{"text": "Submit", "exact": false}'),
        ("click_element",    "Click by CSS selector",                   '{"selector": "#btn-save"}'),
        ("fill_input",       "Fill input by placeholder",               '{"placeholder": "Search...", "value": "test"}'),
        ("fill_input",       "Fill input by label text",                '{"label": "Description", "value": "my text"}'),
        ("get_adjacent_text","Capture text next to a label",            '{"label": "Status"}'),
        ("click_row_icon",   "Click icon in table row (by index)",      '{"search_text": "Row A", "icon_index": 0}'),
        ("click_row_icon",   "Click icon in table row (by name)",       '{"search_text": "Row A", "icon_name": "edit"}'),
        ("search_input",     "Type in search/combo box",                '{"placeholder": "Find...", "value": "Acme"}'),
        ("check_url",        "Assert URL contains text",                '{"expected": "/dashboard"}'),
        ("check_visible",    "Assert element is visible",               '{"selector": "h1"}'),
        ("check_text",       "Assert element text contains value",      '{"selector": ".status", "expected": "Active"}'),
        ("wait_for",         "Wait for element to appear",              '{"selector": ".spinner", "timeout": 5000}'),
        ("wait_for_url",     "Wait for URL to match glob",              '{"pattern": "**/home**"}'),
        ("demo_mouse_move",  "Sweep cursor across page (demo effect)",  '{"steps": 10, "duration_ms": 1000}'),
        ("", "", ""),
        ("ENV VARS", "", ""),
        ("${APP_URL}",       "Application URL from .env",               ""),
        ("${SSO_EMAIL}",     "SSO email from .env",                     ""),
        ("${ANY_VAR}",       "Any environment variable",                ""),
    ]
    info.column_dimensions["A"].width = 22
    info.column_dimensions["B"].width = 48
    info.column_dimensions["C"].width = 42
    hf = Font(bold=True, color="FFFFFF")
    hfi = PatternFill("solid", fgColor="2E4057")
    for r, row in enumerate(instructions, start=1):
        for c, val in enumerate(row, start=1):
            cell = info.cell(row=r, column=c, value=val)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if r == 1 or (not row[1] and not row[2] and row[0]):
                cell.font = hf
                cell.fill = hfi

    wb.save(path)
    print(f"  ✓ Sample Excel created → {path}")
    print(f"    Edit the 'Test Cases' sheet, then run:")
    print(f"    python runner.py")


if __name__ == "__main__":
    out = Path(sys.argv[1]) if len(sys.argv) > 1 else OUTPUT
    create_excel(out)
